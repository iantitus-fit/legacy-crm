"""CSV/XLSX contact importer.

Two-step flow: preview (parses file, suggests mappings, returns stats and a
preview_token), then confirm (uses the cached token + user-confirmed
mappings to bulk-insert contacts). The preview cache is a process-local
in-memory dict — fine for single-instance deployments, would need Redis
or similar to scale horizontally.
"""
from __future__ import annotations

import csv
import io
import re
import threading
import time
import uuid
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

from sqlalchemy.orm import Session

from app.models.contact import Contact
from app.models.contact_import import ContactImport


# ---------- Preview cache ----------

_PREVIEW_CACHE: Dict[str, Dict[str, Any]] = {}
_PREVIEW_LOCK = threading.Lock()
_PREVIEW_TTL_SECONDS = 60 * 30  # 30 minutes


def _purge_expired() -> None:
    now = time.time()
    expired = [
        token
        for token, payload in _PREVIEW_CACHE.items()
        if payload.get("_expires_at", 0) < now
    ]
    for token in expired:
        _PREVIEW_CACHE.pop(token, None)


def _store_preview(payload: Dict[str, Any]) -> str:
    token = uuid.uuid4().hex
    with _PREVIEW_LOCK:
        _purge_expired()
        _PREVIEW_CACHE[token] = {
            **payload,
            "_expires_at": time.time() + _PREVIEW_TTL_SECONDS,
        }
    return token


def get_preview(token: str) -> Optional[Dict[str, Any]]:
    with _PREVIEW_LOCK:
        _purge_expired()
        return _PREVIEW_CACHE.get(token)


def clear_preview(token: str) -> None:
    with _PREVIEW_LOCK:
        _PREVIEW_CACHE.pop(token, None)


# ---------- Parsing ----------

MAX_FILE_BYTES = 50 * 1024 * 1024  # 50 MB
MAX_PREVIEW_SAMPLE = 25
SUPPORTED_EXTS = (".csv", ".xlsx", ".xls")


class ImportError(Exception):
    """Raised for user-correctable parse/format problems. Carries an
    HTTP status code that the router maps to the response."""

    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


def _detect_extension(file_name: str) -> str:
    name = (file_name or "").lower()
    for ext in SUPPORTED_EXTS:
        if name.endswith(ext):
            return ext
    return ""


def parse_file(file_name: str, raw: bytes) -> Tuple[List[str], List[Dict[str, Any]]]:
    """Return (column headers, list of row dicts). Raises ImportError on
    bad input."""
    if len(raw) > MAX_FILE_BYTES:
        raise ImportError("File exceeds 50MB limit", status_code=413)

    ext = _detect_extension(file_name)
    if ext == ".csv":
        return _parse_csv(raw)
    if ext in (".xlsx", ".xls"):
        return _parse_xlsx(raw)
    raise ImportError(
        "Unsupported file format. Use CSV or XLSX.", status_code=400
    )


def _parse_csv(raw: bytes) -> Tuple[List[str], List[Dict[str, Any]]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw.decode("latin-1")
        except UnicodeDecodeError as exc:
            raise ImportError(f"Could not decode file: {exc}") from exc

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(text), dialect=dialect)
    try:
        header = next(reader)
    except StopIteration as exc:
        raise ImportError("File is empty") from exc

    columns = [h.strip() for h in header]
    if not any(columns):
        raise ImportError("Header row is empty")

    rows: List[Dict[str, Any]] = []
    for raw_row in reader:
        if not any((cell or "").strip() for cell in raw_row):
            continue
        row = {
            columns[i]: (raw_row[i].strip() if i < len(raw_row) else "")
            for i in range(len(columns))
        }
        rows.append(row)

    return columns, rows


def _parse_xlsx(raw: bytes) -> Tuple[List[str], List[Dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dep is in requirements
        raise ImportError(
            "XLSX support unavailable on the server", status_code=500
        ) from exc

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportError(f"Could not read XLSX: {exc}") from exc

    ws = wb.active
    if ws is None:
        raise ImportError("XLSX has no active sheet")

    iterator = ws.iter_rows(values_only=True)
    try:
        header_row = next(iterator)
    except StopIteration as exc:
        raise ImportError("File is empty") from exc

    columns = [str(c).strip() if c is not None else "" for c in header_row]
    columns = [c for c in columns if c]
    if not columns:
        raise ImportError("Header row is empty")

    rows: List[Dict[str, Any]] = []
    for raw_row in iterator:
        if raw_row is None:
            continue
        cleaned = [
            ("" if c is None else str(c).strip()) for c in raw_row
        ]
        if not any(cleaned):
            continue
        row = {
            columns[i]: (cleaned[i] if i < len(cleaned) else "")
            for i in range(len(columns))
        }
        rows.append(row)

    return columns, rows


# ---------- Fuzzy mapping ----------

# Maps a normalized column header to a Contact field key. The order matters
# only for the "first match wins" path inside suggest_mappings.
_FIELD_PATTERNS: List[Tuple[str, List[str]]] = [
    ("first_name", [r"^first[\s_-]*name$", r"^fname$", r"^given[\s_-]*name$"]),
    ("last_name", [r"^last[\s_-]*name$", r"^lname$", r"^surname$", r"^family[\s_-]*name$"]),
    (
        "name",
        [
            r"^name$",
            r"^full[\s_-]*name$",
            r"^contact[\s_-]*name$",
            r"^customer[\s_-]*name$",
            r"^client[\s_-]*name$",
            r"primary contact: name$",
        ],
    ),
    (
        "email",
        [
            r"^e[\s_-]*mail$",
            r"^email$",
            r"^email[\s_-]*address$",
            r"primary contact: email$",
        ],
    ),
    (
        "phone",
        [
            r"^phone$",
            r"^phone[\s_-]*number$",
            r"^mobile([\s_-]*phone)?$",
            r"^cell([\s_-]*phone)?$",
            r"^home[\s_-]*phone$",
            r"primary contact: phone$",
        ],
    ),
    (
        "company",
        [
            r"^company([\s_-]*name)?$",
            r"^business([\s_-]*name)?$",
            r"^account([\s_-]*name)?$",
            r"^organization$",
        ],
    ),
    (
        "address",
        [
            r"^address$",
            r"^street$",
            r"^street[\s_-]*address$",
            r"^address[\s_-]*line[\s_-]*1$",
            r"^mailing[\s_-]*street$",
        ],
    ),
    ("city", [r"^city$", r"^mailing[\s_-]*city$"]),
    ("state", [r"^state$", r"^st$", r"^mailing[\s_-]*state$"]),
    (
        "zip_code",
        [
            r"^zip$",
            r"^zip[\s_-]*code$",
            r"^zipcode$",
            r"^postal[\s_-]*code$",
            r"^mailing[\s_-]*zip$",
        ],
    ),
    (
        "lead_source",
        [r"^source$", r"^lead[\s_-]*source$"],
    ),
    (
        "client_type",
        [r"^type$", r"^client[\s_-]*type$", r"^contact[\s_-]*type$"],
    ),
    (
        "notes",
        [r"^notes?$", r"^comments?$", r"^description$"],
    ),
]

VALID_TARGETS = {
    "name",
    "first_name",
    "last_name",
    "email",
    "phone",
    "company",
    "address",
    "city",
    "state",
    "zip_code",
    "lead_source",
    "client_type",
    "notes",
}


def _normalize_header(header: str) -> str:
    return re.sub(r"\s+", " ", header.strip().lower())


def suggest_mappings(columns: List[str]) -> Dict[str, str]:
    """Return {column_header: target_field} for columns that match a
    known field. Columns that don't match are omitted."""
    suggestions: Dict[str, str] = {}
    used_targets: set = set()
    for col in columns:
        norm = _normalize_header(col)
        for target, patterns in _FIELD_PATTERNS:
            if target in used_targets and target not in ("notes",):
                continue
            for pat in patterns:
                if re.search(pat, norm):
                    suggestions[col] = target
                    used_targets.add(target)
                    break
            if col in suggestions:
                break
    return suggestions


# ---------- Preview ----------

def build_preview(
    *,
    file_name: str,
    columns: List[str],
    rows: List[Dict[str, Any]],
    db: Session,
) -> Dict[str, Any]:
    """Compute preview stats + suggested mappings. Caches parsed rows
    against a returned token."""
    if not rows:
        raise ImportError("No data rows found in file")

    suggested = suggest_mappings(columns)
    name_target_columns = {
        col for col, target in suggested.items() if target == "name"
    }
    has_first_last = (
        any(t == "first_name" for t in suggested.values())
        and any(t == "last_name" for t in suggested.values())
    )
    has_name_column = bool(name_target_columns) or has_first_last

    # Existing contacts for duplicate detection
    existing = (
        db.query(Contact.email, Contact.phone, Contact.name, Contact.address)
        .filter(Contact.deleted_at.is_(None))
        .all()
    )
    existing_emails = {
        (e or "").strip().lower() for e, _, _, _ in existing if e
    }
    existing_phones = {
        _digits(p) for _, p, _, _ in existing if p and _digits(p)
    }
    existing_name_addr = {
        ((n or "").strip().lower(), (a or "").strip().lower())
        for _, _, n, a in existing
        if n
    }

    skip_reasons = {"no_name": 0, "empty_row": 0}
    importable = 0
    duplicate_count = 0

    for row in rows:
        if not any((v or "").strip() for v in row.values()):
            skip_reasons["empty_row"] += 1
            continue
        candidate_name = _row_to_name(row, suggested)
        if not candidate_name:
            skip_reasons["no_name"] += 1
            continue
        importable += 1
        candidate_email = _row_field(row, suggested, "email").lower()
        candidate_phone_digits = _digits(_row_field(row, suggested, "phone"))
        candidate_address = _row_field(row, suggested, "address").lower()
        is_dup = (
            (candidate_email and candidate_email in existing_emails)
            or (
                candidate_phone_digits
                and candidate_phone_digits in existing_phones
            )
            or (
                (candidate_name.lower(), candidate_address)
                in existing_name_addr
            )
        )
        if is_dup:
            duplicate_count += 1

    skipped_rows = sum(skip_reasons.values())
    sample_rows = rows[:MAX_PREVIEW_SAMPLE]

    unmapped_columns = [c for c in columns if c not in suggested]

    payload = {
        "file_name": file_name,
        "total_rows": len(rows),
        "importable_rows": importable,
        "skipped_rows": skipped_rows,
        "skip_reasons": skip_reasons,
        "potential_duplicates": duplicate_count,
        "columns_detected": columns,
        "suggested_mappings": suggested,
        "unmapped_columns": unmapped_columns,
        "sample_rows": sample_rows,
        "warnings": [] if has_name_column else ["No name column detected."],
        # Internal — not in response
        "_columns": columns,
        "_rows": rows,
    }
    token = _store_preview(payload)
    payload["preview_token"] = token
    return payload


# ---------- Confirm / import ----------

def _digits(value: str) -> str:
    return re.sub(r"\D", "", value or "")


_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _is_valid_email(value: str) -> bool:
    return bool(value) and bool(_EMAIL_RE.match(value))


def _row_field(
    row: Dict[str, Any], mappings: Dict[str, str], target: str
) -> str:
    """First non-empty value among columns mapped to target."""
    for col, mapped in mappings.items():
        if mapped == target:
            v = (row.get(col) or "").strip()
            if v:
                return v
    return ""


def _row_to_name(row: Dict[str, Any], mappings: Dict[str, str]) -> str:
    direct = _row_field(row, mappings, "name")
    if direct:
        return direct
    first = _row_field(row, mappings, "first_name")
    last = _row_field(row, mappings, "last_name")
    combined = " ".join(part for part in (first, last) if part).strip()
    return combined


def _build_notes(
    row: Dict[str, Any],
    mappings: Dict[str, str],
    combine_unmapped: bool,
    columns: List[str],
) -> str:
    parts: List[str] = []
    direct_notes = _row_field(row, mappings, "notes")
    if direct_notes:
        parts.append(direct_notes)
    if combine_unmapped:
        for col in columns:
            mapped = mappings.get(col)
            if mapped:
                continue
            v = (row.get(col) or "").strip()
            if v:
                parts.append(f"{col}: {v}")
    return "\n".join(parts)


def confirm_import(
    *,
    db: Session,
    user_id: Optional[int],
    preview_token: str,
    field_mappings: Dict[str, Optional[str]],
    options: Dict[str, Any],
) -> ContactImport:
    cached = get_preview(preview_token)
    if not cached:
        raise ImportError("Preview token expired or invalid", status_code=400)

    columns: List[str] = cached["_columns"]
    rows: List[Dict[str, Any]] = cached["_rows"]
    file_name: str = cached["file_name"]

    # Sanitize mappings: ignore unknown columns, validate target values
    sanitized: Dict[str, str] = {}
    for col, target in (field_mappings or {}).items():
        if col not in columns or target is None:
            continue
        if target not in VALID_TARGETS:
            continue
        sanitized[col] = target

    default_client_type = options.get("default_client_type")
    default_lead_source = options.get("default_lead_source") or "CSV Import"
    duplicate_handling = options.get("duplicate_handling") or "skip"
    combine_unmapped = bool(options.get("combine_unmapped_to_notes", True))
    pipeline_id = options.get("pipeline_id")
    stage_id = options.get("stage_id")

    has_client_type_mapping = any(
        t == "client_type" for t in sanitized.values()
    )
    has_lead_source_mapping = any(
        t == "lead_source" for t in sanitized.values()
    )

    # Existing duplicates
    existing = (
        db.query(Contact.email, Contact.phone, Contact.name, Contact.address)
        .filter(Contact.deleted_at.is_(None))
        .all()
    )
    existing_emails = {
        (e or "").strip().lower() for e, _, _, _ in existing if e
    }
    existing_phones = {
        _digits(p) for _, p, _, _ in existing if p and _digits(p)
    }
    existing_name_addr = {
        ((n or "").strip().lower(), (a or "").strip().lower())
        for _, _, n, a in existing
        if n
    }

    import_id = uuid.uuid4().hex
    started = time.time()

    imported_count = 0
    skipped_invalid = 0
    skipped_duplicates = 0
    duplicate_total = 0

    new_contacts: List[Contact] = []
    BATCH_SIZE = 500

    for row in rows:
        if not any((v or "").strip() for v in row.values()):
            skipped_invalid += 1
            continue
        name = _row_to_name(row, sanitized)
        if not name:
            skipped_invalid += 1
            continue

        email_raw = _row_field(row, sanitized, "email")
        email = email_raw if _is_valid_email(email_raw) else None
        phone = _row_field(row, sanitized, "phone") or None
        company = _row_field(row, sanitized, "company") or None
        address = _row_field(row, sanitized, "address") or None
        city = _row_field(row, sanitized, "city") or None
        state = _row_field(row, sanitized, "state") or "IN"
        zip_code = _row_field(row, sanitized, "zip_code") or None

        lead_source = _row_field(row, sanitized, "lead_source")
        if not lead_source:
            lead_source = default_lead_source

        client_type = _row_field(row, sanitized, "client_type")
        if not client_type and default_client_type:
            client_type = default_client_type

        notes_text = _build_notes(row, sanitized, combine_unmapped, columns)

        # Duplicate check
        is_dup = False
        email_lc = (email or "").lower()
        phone_digits = _digits(phone or "")
        addr_lc = (address or "").lower()
        if email_lc and email_lc in existing_emails:
            is_dup = True
        elif phone_digits and phone_digits in existing_phones:
            is_dup = True
        elif (name.lower(), addr_lc) in existing_name_addr:
            is_dup = True

        if is_dup:
            duplicate_total += 1
            if duplicate_handling == "skip":
                skipped_duplicates += 1
                continue

        def _trunc(value: Optional[str], n: int) -> Optional[str]:
            return value[:n] if value else None

        contact = Contact(
            name=name[:255],
            email=email,
            phone=_trunc(phone, 50),
            company=_trunc(company, 255),
            address=_trunc(address, 500),
            city=_trunc(city, 100),
            state=(state or "IN")[:2],
            zip=_trunc(zip_code, 10),
            lead_source=_trunc(lead_source, 100),
            client_type=_trunc(client_type, 20),
            pipeline_id=pipeline_id,
            stage_id=stage_id,
            import_id=import_id,
            import_source_file=(file_name[:500] if file_name else None),
        )
        new_contacts.append(contact)

        # Track within-batch dedup
        if email_lc:
            existing_emails.add(email_lc)
        if phone_digits:
            existing_phones.add(phone_digits)
        existing_name_addr.add((name.lower(), addr_lc))

        imported_count += 1

        if notes_text:
            # Stash notes on the contact instance temporarily — applied
            # after flush so we have the contact ID.
            contact._pending_note = notes_text  # type: ignore[attr-defined]

        if len(new_contacts) >= BATCH_SIZE:
            _flush_batch(db, new_contacts, user_id)
            new_contacts.clear()

    if new_contacts:
        _flush_batch(db, new_contacts, user_id)
        new_contacts.clear()

    duration = round(time.time() - started, 2)

    record = ContactImport(
        id=import_id,
        user_id=user_id,
        file_name=file_name,
        total_rows=len(rows),
        imported_count=imported_count,
        skipped_count=skipped_invalid,
        duplicate_count=duplicate_total,
        field_mappings=sanitized,
        options={
            "default_client_type": default_client_type,
            "default_lead_source": default_lead_source,
            "duplicate_handling": duplicate_handling,
            "combine_unmapped_to_notes": combine_unmapped,
            "pipeline_id": pipeline_id,
            "stage_id": stage_id,
        },
        duration_seconds=duration,
    )
    db.add(record)
    db.commit()
    db.refresh(record)

    # Transient: distinguishes "duplicates we skipped" from total duplicates
    record._skipped_duplicates = skipped_duplicates  # type: ignore[attr-defined]

    clear_preview(preview_token)
    return record


def _flush_batch(
    db: Session, contacts: List[Contact], user_id: Optional[int]
) -> None:
    pending_notes: List[Tuple[Contact, str]] = [
        (c, getattr(c, "_pending_note", "")) for c in contacts
    ]
    db.add_all(contacts)
    db.flush()
    # Materialize notes via the Note model for any contacts with notes.
    # Note.created_by_user_id is NOT NULL — without a user, drop the note.
    if user_id is None:
        return
    from app.models.note import Note

    for contact, note_text in pending_notes:
        if not note_text:
            continue
        db.add(
            Note(
                entity_type="contact",
                entity_id=contact.id,
                note_type="company",
                content=note_text,
                created_by_user_id=user_id,
            )
        )
    db.flush()


# ---------- Undo ----------

def undo_import(db: Session, import_id: str) -> int:
    """Soft-delete all contacts created by an import. Returns count."""
    record = (
        db.query(ContactImport).filter(ContactImport.id == import_id).first()
    )
    if not record:
        raise ImportError("Import not found", status_code=404)
    if record.undone_at is not None:
        raise ImportError("Import already undone", status_code=400)

    now = datetime.now(timezone.utc)
    affected = (
        db.query(Contact)
        .filter(Contact.import_id == import_id, Contact.deleted_at.is_(None))
        .update({Contact.deleted_at: now}, synchronize_session=False)
    )
    record.undone_at = now
    db.commit()
    return int(affected)
